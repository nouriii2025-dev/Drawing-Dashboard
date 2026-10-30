"""Import the legacy Excel workbook into the database.

    python manage.py import_excel data/Daily_Update_27-10-2026.xlsx --replace

What it handles:
  * "Daily Update" sheet - every column is kept. Month header rows (DEC, JAN, ...)
    are not records; they are used to repair dates.
  * Excel read many typed dd/mm dates as mm/dd (e.g. 12 Feb 2025 became 2 Dec 2025).
    The month header the row sits under tells us which reading is correct.
  * Vertically merged cells are filled down so every row is self-contained.
  * "Factory Drawing" and "Projects" sheets become their own tables.
"""
import datetime as dt
import re

from django.core.management.base import BaseCommand
from django.db import transaction
from openpyxl import load_workbook

from tracker.models import DailyUpdate, FactoryDrawing, Project

MONTH_NAMES = {"JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUNE", "JUL", "JULY",
               "AUG", "SEP", "SEPT", "OCT", "NOV", "DEC"}
MONTH_NUM = {"JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MAY": 5, "JUN": 6, "JUNE": 6, "JUL": 7,
             "JULY": 7, "AUG": 8, "SEP": 9, "SEPT": 9, "OCT": 10, "NOV": 11, "DEC": 12}
DATE_RE = re.compile(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})$")
DWG_NORM_RE = re.compile(
    r"^(?:ACCU[-\s.]*)?(\d{3,5})\s*([A-Za-z])?\s*$",
    re.IGNORECASE,
)


def _normalize_dwg(value: str) -> str:
    """Force ACCU-XXXX (optional letter suffix) when the value is a plain number or Accu-/ACCU- form."""
    if not value:
        return value
    s = value.strip()
    if re.fullmatch(r"ACCU-\d{3,5}[A-Z]?", s):
        return s
    m = DWG_NORM_RE.match(s)
    if not m:
        return value
    num, letter = m.group(1), (m.group(2) or "").upper()
    return f"ACCU-{num}{letter}"


# Daily Update column index (1-based) -> model field
COLS = {
    1: "sno", 2: "date", 3: "dwg_no", 4: "drawn_by", 5: "software", 6: "item", 7: "model_no",
    8: "customer", 9: "project_po", 10: "sq_number", 11: "note", 12: "status",
    13: "sales_person", 14: "approved_by", 15: "submitted_date", 16: "issued_for",
    17: "dwg_status", 18: "factory_dwg_status", 19: "file_path", 20: "price", 21: "remarks",
}


def text(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(round(v, 6))
    if isinstance(v, dt.datetime):
        return v.strftime("%d-%m-%Y")
    return re.sub(r"[ \t]+", " ", str(v).replace("\xa0", " ")).strip()


def parse_text_date(s):
    """'14/01/2026', '15-01-2025', '21.08.2026' -> (date | None, matched?)"""
    m = DATE_RE.match(s.strip())
    if not m:
        return None, False
    d, mo, y = map(int, m.groups())
    try:
        return dt.date(y, mo, d), True
    except ValueError:
        return None, True


class Command(BaseCommand):
    help = "Import the Daily Update Excel workbook."

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--replace", action="store_true", help="Delete existing records first.")

    @transaction.atomic
    def handle(self, path, replace=False, **opts):
        wb = load_workbook(path, data_only=True)
        if replace:
            DailyUpdate.objects.all().delete()
            FactoryDrawing.objects.all().delete()
            Project.objects.all().delete()
        n = self.import_updates(wb["Daily Update"])
        f = self.import_factory(wb["Factory Drawing"]) if "Factory Drawing" in wb.sheetnames else 0
        p = self.import_projects(wb["Projects"]) if "Projects" in wb.sheetnames else 0
        self.stdout.write(self.style.SUCCESS(
            f"Imported {n} daily updates, {f} factory drawings, {p} projects."))

    # ------------------------------------------------------------------
    def import_updates(self, ws):
        fill = {}
        for rng in ws.merged_cells.ranges:
            if rng.min_row == rng.max_row:
                continue
            top = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row + 1, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    fill[(r, c)] = top

        def raw(r, c):
            return ws.cell(r, c).value

        def val(r, c):
            v = raw(r, c)
            return v if v is not None else fill.get((r, c))

        section = None            # (year, month) from month header rows
        last_period = (2024, 8)   # before the first header the sheet starts in Aug 2024
        parent_date = None
        imported, mismatches = 0, []

        for r in range(2, ws.max_row + 1):
            b = raw(r, 2)
            others = [text(raw(r, c)) for c in range(3, 22) if c != 4]
            blank_row = all(o in ("", "-") for o in others)
            # ---- month / year header rows
            if isinstance(b, str) and b.strip().upper() in MONTH_NAMES and blank_row:
                name_month = MONTH_NUM[b.strip().upper()]
                if section is None:
                    section = (2024, name_month)
                else:
                    y, m = section
                    section = (y + 1, 1) if m == 12 else (y, m + 1)  # sequence beats typos ("AUG" for Oct)
                a = raw(r, 1)
                if isinstance(a, int) and a >= 2000:
                    section = (a, section[1])  # explicit year marker, e.g. "2026" beside JAN
                last_period = section
                continue
            if isinstance(b, (int, float)) and blank_row:
                continue  # lone year marker such as "2025"

            sno_raw = raw(r, 1)
            continuation = sno_raw in (None, "")
            rec = {}
            for c, field in COLS.items():
                if field in ("date", "submitted_date", "sno"):
                    continue
                rec[field] = text(val(r, c))
            rec["status"] = rec["status"].strip()

            s = val(r, 1)
            rec["sno"] = int(s) if isinstance(s, (int, float)) and not isinstance(s, bool) else None

            # ---- Date column
            rec["date"], rec["date_text"] = None, ""
            d = raw(r, 2)
            section_month = section[1] if section else None
            if isinstance(d, dt.datetime):
                d0 = d.date()
                cands = [d0]
                if d0.day <= 12 and d0.day != d0.month:
                    cands.append(dt.date(d0.year, d0.day, d0.month))  # day/month swapped by Excel
                chosen = d0
                if continuation and parent_date:
                    later = sorted(c for c in cands if c >= parent_date)
                    chosen = later[0] if later else d0
                elif section_month:
                    match = [c for c in cands if c.month == section_month]
                    chosen = match[0] if match else d0
                rec["date"] = chosen
            elif isinstance(d, str) and d.strip():
                parsed, matched = parse_text_date(d)
                if parsed:
                    rec["date"] = parsed
                elif matched or not re.search(r"[A-Za-z]", d):
                    rec["date_text"] = d.strip()  # impossible date such as 31/09/2025
                else:  # a comment typed into the Date column
                    extra = d.strip()
                    rec["note"] = extra if rec["note"] in ("", "-") else f"{extra} - {rec['note']}"

            # ---- Submitted date
            sd = raw(r, 15)
            rec["submitted_date"] = None
            if isinstance(sd, dt.datetime):
                rec["submitted_date"] = sd.date()
            elif isinstance(sd, str) and sd.strip():
                rec["submitted_date"], _ = parse_text_date(sd)

            # ---- skip rows with nothing meaningful (blank spacer rows, stray 'Aravindh')
            core = [rec[k] for k in rec if k not in ("drawn_by", "date", "submitted_date", "sno")]
            if not any(core) and not rec["date"] and not rec["submitted_date"] and rec["sno"] is None:
                continue

            if rec["date"]:
                period = (rec["date"].year, rec["date"].month)
                if not continuation:
                    parent_date = rec["date"]
                    if section and period != section:
                        mismatches.append((r, rec["date"], section))
            else:
                period = last_period
            rec["period"] = f"{period[0]}-{period[1]:02d}"
            if section is None and rec["date"]:
                last_period = period

            # Normalize DWG No to ACCU-XXXX when it is a plain number or Accu-/ACCU- form
            rec["dwg_no"] = _normalize_dwg(rec.get("dwg_no", ""))

            imported += 1
            DailyUpdate.objects.create(position=imported, **rec)

        for r, d, sec in mismatches:
            self.stdout.write(f"  note: row {r} date {d} sits under {sec[0]}-{sec[1]:02d}")
        return imported

    # ------------------------------------------------------------------
    def import_factory(self, ws):
        labels = {"location": "location", "item": "item", "factory name": "factory_name",
                  "enduser": "end_user", "forwaed by": "forwarded_by", "subjuct": "subject",
                  "forward me": "forwarded_to", "recived date": "received_date"}
        records, cur = [], None
        for r in range(1, ws.max_row + 1):
            b = text(ws.cell(r, 2).value)
            if not b:
                continue
            key = labels.get(b.rstrip(" :").lower())
            if key == "location":
                cur = {"location": text(ws.cell(r, 3).value)}
                records.append(cur)
            elif cur is None:
                continue
            elif key == "received_date":
                parsed, _ = parse_text_date(text(ws.cell(r, 3).value))
                cur["received_date"] = parsed
            elif key:
                cur[key] = text(ws.cell(r, 3).value)
                if key == "item":
                    cur["spec"] = text(ws.cell(r, 5).value)
            else:
                cur["note"] = b
        for i, rec in enumerate(records, 1):
            FactoryDrawing.objects.create(position=i, **rec)
        return len(records)

    def import_projects(self, ws):
        client, po = text(ws["A2"].value), text(ws["B2"].value)
        sections = "\n".join(t for t in (text(ws["A4"].value), text(ws["A5"].value)) if t)
        if not (client or po):
            return 0
        Project.objects.create(position=1, client=client, project_po=po, sections=sections)
        return 1
